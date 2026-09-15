#!/usr/bin/env python3
"""Generate docs/realdreamteam-overview.pptx (3 slides) with python-pptx.

Usage:  pip install python-pptx && python3 docs/generate_overview_deck.py

Images: reuses docs/mockup-summary.png and docs/architecture.png. Re-render them with
  npx playwright screenshot --full-page --viewport-size=1200,800 docs/mockup-summary.html docs/mockup-summary.png
  npx -p @mermaid-js/mermaid-cli mmdc -i docs/architecture.mmd -o docs/architecture.png -b white -w 1400
"""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

DOCS = Path(__file__).resolve().parent
OUT = DOCS / "realdreamteam-overview.pptx"
MOCKUP = DOCS / "mockup-summary.png"
ARCH = DOCS / "architecture.png"

NAVY = RGBColor(0x1F, 0x2A, 0x44)
INK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x5F, 0x63, 0x68)
ACCENT = RGBColor(0xB4, 0x5F, 0x06)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.5)


def add_text(slide, left, top, width, height, paragraphs, size=14, color=INK,
             bold=False, anchor=MSO_ANCHOR.TOP):
    """paragraphs: list of str or (str, {opts}) where opts may set size/bold/color/bullet/space_after."""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    for i, item in enumerate(paragraphs):
        text, opts = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(opts.get("space_after", 6))
        if opts.get("bullet", False):
            text = "\u2022  " + text
        r = p.add_run()
        r.text = text
        r.font.name = FONT
        r.font.size = Pt(opts.get("size", size))
        r.font.bold = opts.get("bold", bold)
        r.font.color.rgb = opts.get("color", color)
    return box


def add_header(slide, title, kicker):
    bar = slide.shapes.add_shape(1, 0, 0, W, Inches(1.15))  # 1 = rectangle
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    add_text(slide, MARGIN, Inches(0.12), W - 2 * MARGIN, Inches(0.35),
             [kicker], size=12, color=RGBColor(0xC9, 0xD3, 0xE6))
    add_text(slide, MARGIN, Inches(0.42), W - 2 * MARGIN, Inches(0.65),
             [title], size=28, color=WHITE, bold=True)


def add_footer(slide, n):
    add_text(slide, MARGIN, H - Inches(0.45), W - 2 * MARGIN, Inches(0.3),
             [f"COG-GTM/realdreamteam  \u00b7  docs/generate_overview_deck.py  \u00b7  {n}/3"],
             size=9, color=GREY)


def add_picture_fit(slide, path, left, top, max_w, max_h):
    """Insert an image scaled to fit inside the box, centred."""
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(max_w / iw, max_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    pic = slide.shapes.add_picture(str(path), left + (max_w - w) // 2, top + (max_h - h) // 2, Emu(w), Emu(h))
    pic.line.color.rgb = RGBColor(0xD0, 0xD4, 0xDA)
    pic.line.width = Pt(0.75)
    return pic


def bullets(items, size=14):
    return [(t, {"bullet": True, "size": size, "space_after": 9}) for t in items]


def slide_overview(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(s, "Real Dream Team \u2014 internal auction app", "1 \u00b7 Overview")
    body_top = Inches(1.4)
    body_h = H - body_top - Inches(0.6)
    img_w = Inches(5.0)
    add_picture_fit(s, MOCKUP, MARGIN, body_top, img_w, body_h)
    add_text(s, MARGIN, H - Inches(0.75), img_w, Inches(0.3),
             ["Summary page (docs/mockup-summary.html)"], size=10, color=GREY)
    left = MARGIN + img_w + Inches(0.4)
    add_text(s, left, body_top, W - left - MARGIN, body_h, [
        ("Live at rdt-auction.marklovestech.com", {"size": 14, "bold": True, "color": ACCENT, "space_after": 10}),
        *bullets([
            "Sign in: site code entry + name picker",
            "Summary: Notifications, \u201cMatches your interests\u201d, and Discover (5 random open lots)",
            "Browse auctions and lots, \u2605 favorite, place whole-number bids that must beat the high bid",
            "Interest matching by category / artist / keyword, with \u201cwhy it matched\u201d chips",
            "Auction close \u2192 SOLD ribbon, hammer price, winner, sound, Won/Lost notifications",
            "/admin (second code): add lots, edit close times, close/reopen auctions, ban users",
        ], size=16),
    ])
    add_footer(s, 1)


def slide_architecture(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(s, "Architecture", "2 \u00b7 Three moving parts, all always on")
    top = Inches(1.35)
    add_picture_fit(s, ARCH, MARGIN, top, W - 2 * MARGIN, Inches(2.5))
    add_text(s, MARGIN, top + Inches(2.55), W - 2 * MARGIN, Inches(0.3),
             ["docs/architecture.mmd, rendered with mermaid-cli"], size=10, color=GREY)
    col_top = top + Inches(3.0)
    col_w = (W - 2 * MARGIN - Inches(0.4)) / 2
    add_text(s, MARGIN, col_top, col_w, Inches(2.6), [
        ("Stack", {"size": 15, "bold": True, "color": ACCENT, "space_after": 8}),
        *bullets([
            "Node.js 20 + Express + EJS + hand-written SQL (pg)",
            "No build step, no frontend framework, no API layer \u2014 server renders HTML",
            "Stateless Node process managed by systemd, with a 5 s poller that closes auctions",
            "All state lives in Postgres; npm run db:reset reloads data/seed/*.json",
        ]),
    ])
    add_text(s, MARGIN + col_w + Inches(0.4), col_top, col_w, Inches(2.6), [
        ("Request flow", {"size": 15, "bold": True, "color": ACCENT, "space_after": 8}),
        *bullets([
            "IONOS DNS \u2192 EC2 instance",
            "Caddy terminates TLS on 80/443 (Let\u2019s Encrypt) \u2192 Node on :3000",
            "Node runs SQL against Supabase Postgres \u2192 returns plain HTML",
            "Browser fetches lot images directly from Wikimedia Commons",
        ]),
    ])
    add_footer(s, 2)


def slide_build(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(s, "Collaborative build: GitHub + Devin sessions", "3 \u00b7 How it was built")
    top = Inches(1.4)
    col_w = (W - 2 * MARGIN - Inches(0.4)) / 2
    add_text(s, MARGIN, top, col_w, Inches(5.3), [
        ("One Devin session per concern", {"size": 15, "bold": True, "color": ACCENT, "space_after": 8}),
        ("Sessions indexed in docs/auction-app-build-design.md \u00a7 \u201cBuild process \u2014 sessions\u201d:",
         {"size": 12, "color": GREY, "space_after": 8}),
        *bullets([
            "Design",
            "Permissions & access",
            "Backend database work",
            "Data Model",
            "RDT Synthetic Data",
            "RDT-Quality",
            "RDT-Security",
            "Links & References index",
        ], size=15),
    ])
    right = MARGIN + col_w + Inches(0.4)
    add_text(s, right, top, col_w, Inches(5.3), [
        ("GitHub workflow", {"size": 15, "bold": True, "color": ACCENT, "space_after": 8}),
        *bullets([
            "Each build step = one PR on a devin/* branch, with feature: / bug: commits and screenshots",
            "PRs opened by devin-ai-integration[bot]; reviewed and merged into main by CognitionMark",
            "Examples: #41 scaffold \u00b7 #46 auctions & bidding \u00b7 #44 admin & close \u00b7 "
            "#47 seed data \u00b7 #56 README & architecture",
        ]),
        ("Secrets stay out of the public repo", {"size": 15, "bold": True, "color": ACCENT, "space_after": 8}),
        *bullets([
            "Hosts, IPs, users and ports live in the private Devin knowledge note \u201cRDT infrastructure access\u201d",
            "Credentials are Devin org secrets: AUCTION_DATABASE_URL, AUCTION_DATABASE_PASSWORD, RDT_EC2_SSH_KEY",
        ]),
    ])
    add_footer(s, 3)


def main():
    for p in (MOCKUP, ARCH):
        if not p.exists():
            raise SystemExit(f"missing {p} \u2014 see the render commands in this file's docstring")
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    slide_overview(prs)
    slide_architecture(prs)
    slide_build(prs)
    prs.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
